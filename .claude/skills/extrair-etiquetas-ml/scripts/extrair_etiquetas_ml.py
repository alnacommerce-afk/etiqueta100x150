#!/usr/bin/env python3
"""Extrai Código ML, quantidade e SKU de um PDF de preparação de envio do Mercado Livre
e grava um .xlsx com exatamente as colunas: Código ML | quantidade | SKU.

Nada é inventado: cada campo vem do PDF, associado ao produto pela posição (coordenadas)
na tabela. Campo que não puder ser confirmado fica em branco e o resultado é marcado
INCOMPLETO (exit code 1). Uso:  python extrair_etiquetas_ml.py ARQUIVO.pdf [-o saida.xlsx]
"""
from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pymupdf

COLUNAS = ["Código ML", "quantidade", "SKU"]
RE_ANCHOR = re.compile(r"^C.digo ML:\s*(\S*)")  # "Código ML: XXXX Código universal: ..."
RE_TOTAIS = re.compile(r"Produtos do envio:\s*(\d+)\s*\|\s*Total de unidades:\s*(\d+)", re.I)
RE_FRETE = re.compile(r"(?:Frete|Envio)\s*#\s*(\d+)", re.I)
RE_CODIGO_ML = re.compile(r"^[A-Z]{4}\d{5}$")
RE_SKU_OK = re.compile(r"^[A-Za-z0-9_.\-]+$")
SKU_LINE_GAP = 9.5   # pt: linha de continuação do SKU fica ~8 pt abaixo; a descrição, ~12 pt
ROW_MARGIN = 12      # pt acima da linha "Código ML" onde a linha da tabela começa


@dataclass
class Produto:
    pagina: int
    posicao: int  # posição global (1-based) na ordem do documento
    codigo_ml: str | None = None
    quantidade: int | None = None
    sku: str | None = None
    problemas: list[str] = field(default_factory=list)

    @property
    def completo(self) -> bool:
        return bool(self.codigo_ml) and self.quantidade is not None and bool(self.sku)


@dataclass
class Resultado:
    pdf: Path
    frete: str | None
    produtos: list[Produto]
    total_produtos_doc: int | None
    total_unidades_doc: int | None
    avisos: list[str] = field(default_factory=list)
    erros: list[str] = field(default_factory=list)  # impedem declarar sucesso

    @property
    def completo(self) -> bool:
        return not self.erros and bool(self.produtos) and all(p.completo for p in self.produtos)


def _linhas(page: pymupdf.Page, tp=None):
    d = page.get_text("dict", textpage=tp) if tp else page.get_text("dict")
    out = []
    for b in d["blocks"]:
        for ln in b.get("lines", []):
            t = "".join(s["text"] for s in ln["spans"]).strip()
            if t:
                out.append((t, ln["bbox"][0], ln["bbox"][1], ln["bbox"][3]))
    out.sort(key=lambda r: (round(r[2], 0), r[1]))
    return out


def _palavras(page: pymupdf.Page, tp=None):
    return page.get_text("words", textpage=tp) if tp else page.get_text("words")


def _ocr(page: pymupdf.Page):
    for lang in ("por+eng", "eng"):
        try:
            return page.get_textpage_ocr(language=lang, dpi=300, full=True)
        except Exception:
            continue
    return None


def _extrair_pagina(page: pymupdf.Page, inicio_pos: int, res: Resultado) -> list[Produto]:
    tp = None
    linhas = _linhas(page)
    if not any(RE_ANCHOR.match(t) for t, *_ in linhas) and len(page.get_text().strip()) < 30:
        tp = _ocr(page)  # página sem texto digital: tenta OCR (requer Tesseract)
        if tp is None:
            res.erros.append(f"Página {page.number + 1}: sem texto digital e OCR indisponível (instale o Tesseract).")
            return []
        res.avisos.append(f"Página {page.number + 1}: lida por OCR; confira visualmente todos os campos.")
        linhas = _linhas(page, tp)
    palavras = _palavras(page, tp)

    anchors = [(i, ln) for i, ln in enumerate(linhas) if RE_ANCHOR.match(ln[0])]
    if not anchors:
        return []

    # coluna UNIDADES: do x do cabeçalho até o início do cabeçalho seguinte (IDENTIFICAÇÃO)
    hdr = sorted((w for w in palavras if w[4] == "UNIDADES"), key=lambda w: w[1])  # cabeçalho em maiúsculas, o mais alto
    col = None
    if hdr:
        h = hdr[0]
        prox = [w for w in palavras if w[4].startswith("IDENTIFICA") and w[0] > h[0] and abs(w[1] - h[1]) < 4]
        col = (h[0] - 12, (min(w[0] for w in prox) if prox else h[0] + 40) - 4, h[3])

    produtos: list[Produto] = []
    for n, (idx, (texto, x0, y0, _y1)) in enumerate(anchors):
        prod = Produto(pagina=page.number + 1, posicao=inicio_pos + n)
        rotulo = f"página {prod.pagina}, produto {prod.posicao}"

        # Código ML
        cod = RE_ANCHOR.match(texto).group(1)
        if cod:
            prod.codigo_ml = cod
            if not RE_CODIGO_ML.match(cod):
                res.avisos.append(f"{rotulo}: Código ML '{cod}' fora do padrão usual (4 letras + 5 dígitos); mantido como está no PDF.")
        else:
            prod.problemas.append("Código ML ausente")

        # SKU: texto após "SKU:" na linha do cabeçalho; se vazio, linha(s) logo abaixo (gap curto)
        m = re.search(r"SKU:\s*(.*)$", texto)
        if not m:
            prod.problemas.append("rótulo 'SKU:' não encontrado")
        else:
            partes = [m.group(1).strip()] if m.group(1).strip() else []
            prev_y = y0
            for t2, x2, ya, _yb in linhas[idx + 1:]:
                if ya - prev_y > SKU_LINE_GAP or RE_ANCHOR.match(t2):
                    break
                if abs(x2 - x0) > 6:  # texto de outra coluna na mesma altura
                    continue
                partes.append(t2)
                prev_y = ya
                if len(partes) >= 2:
                    break
            sku = "".join(partes)
            if not sku:
                prod.problemas.append("SKU ausente")
            else:
                prod.sku = sku
                if len(partes) > 1:
                    res.avisos.append(f"{rotulo}: SKU ocupa {len(partes)} linhas e foi unido sem espaços ('{sku}'); confira.")
                if not RE_SKU_OK.match(sku):
                    res.avisos.append(f"{rotulo}: SKU '{sku}' contém caracteres incomuns; confira.")

        # Quantidade: único inteiro na coluna UNIDADES dentro da faixa vertical da linha da tabela
        topo = y0 - ROW_MARGIN
        base = (anchors[n + 1][1][2] - ROW_MARGIN) if n + 1 < len(anchors) else page.rect.height
        if col is None:
            prod.problemas.append("coluna UNIDADES não encontrada")
        else:
            cands = [w for w in palavras if col[0] <= w[0] <= col[1] and topo <= w[1] < base and w[1] > col[2] - 1]
            if len(cands) == 1 and cands[0][4].isdigit():
                prod.quantidade = int(cands[0][4])
                if prod.quantidade == 0:
                    res.avisos.append(f"{rotulo}: quantidade 0 no PDF; confira.")
            elif not cands:
                prod.problemas.append("quantidade ausente")
            else:
                prod.problemas.append("quantidade ambígua (" + ", ".join(w[4] for w in cands) + ")")

        for pr in prod.problemas:
            res.erros.append(rotulo + (f" ({prod.codigo_ml})" if prod.codigo_ml else "") + f": {pr}")
        produtos.append(prod)
    return produtos


def extrair(pdf: Path) -> Resultado:
    res = Resultado(pdf=pdf, frete=None, produtos=[], total_produtos_doc=None, total_unidades_doc=None)
    doc = pymupdf.open(pdf)
    texto_todo = "\n".join(p.get_text() for p in doc)
    if (m := RE_FRETE.search(texto_todo)):
        res.frete = m.group(1)
    if (m := RE_TOTAIS.search(texto_todo)):
        res.total_produtos_doc, res.total_unidades_doc = int(m.group(1)), int(m.group(2))
    else:
        res.avisos.append("O PDF não informa 'Produtos do envio / Total de unidades' (ou não foi legível): totais não validados.")
    for page in doc:
        res.produtos += _extrair_pagina(page, len(res.produtos) + 1, res)

    if not res.produtos:
        res.erros.append("Nenhum produto encontrado no PDF.")
    if res.total_produtos_doc is not None and res.total_produtos_doc != len(res.produtos):
        res.erros.append(f"Divergência: o PDF informa {res.total_produtos_doc} produtos, mas foram extraídos {len(res.produtos)}.")
    soma = sum(p.quantidade or 0 for p in res.produtos)
    if res.total_unidades_doc is not None and res.total_unidades_doc != soma:
        res.erros.append(f"Divergência: o PDF informa {res.total_unidades_doc} unidades, mas a soma das quantidades extraídas é {soma}.")
    return res


def gravar_xlsx(res: Resultado, destino: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.append(COLUNAS)
    for p in res.produtos:
        ws.append([p.codigo_ml, p.quantidade, p.sku])
    for row in ws.iter_rows(min_row=2):
        for c in (row[0], row[2]):
            if c.value is not None:
                c.data_type = "s"
            c.number_format = "@"  # texto: evita notação científica / conversões
        row[1].number_format = "0"
    for letra, w in zip("ABC", (16, 12, 24)):
        ws.column_dimensions[letra].width = w
    wb.save(destino)


def verificar_xlsx(res: Resultado, caminho: Path) -> list[str]:
    """Reabre o arquivo gravado e compara com o extraído. Retorna as divergências."""
    falhas: list[str] = []
    wb = openpyxl.load_workbook(caminho)
    if len(wb.sheetnames) != 1:
        falhas.append(f"esperada 1 planilha, encontradas {len(wb.sheetnames)}")
    ws = wb.worksheets[0]
    cab = [c.value for c in ws[1]]
    if cab != COLUNAS or ws.max_column != 3:
        falhas.append(f"cabeçalho {cab} diferente de {COLUNAS}")
    linhas = [[c.value for c in r] for r in ws.iter_rows(min_row=2, max_col=3)]
    if len(linhas) != len(res.produtos):
        falhas.append(f"{len(linhas)} linhas no Excel vs {len(res.produtos)} produtos")
    for i, (lin, p) in enumerate(zip(linhas, res.produtos), start=2):
        if lin != [p.codigo_ml, p.quantidade, p.sku]:
            falhas.append(f"linha {i}: {lin} != {[p.codigo_ml, p.quantidade, p.sku]}")
        if lin[1] is not None and (not isinstance(lin[1], int) or isinstance(lin[1], bool)):
            falhas.append(f"linha {i}: quantidade não é inteiro ({type(lin[1]).__name__})")
        for v in (lin[0], lin[2]):
            if v is not None and not isinstance(v, str):
                falhas.append(f"linha {i}: código/SKU não é texto")
    return falhas


def main(argv: list[str] | None = None) -> int:
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8")
        except Exception:
            pass
    ap = argparse.ArgumentParser(description="PDF de preparação do Mercado Livre -> Excel (Código ML, quantidade, SKU)")
    ap.add_argument("pdf", type=Path)
    ap.add_argument("-o", "--saida", type=Path, help="caminho do .xlsx (padrão: ao lado do PDF)")
    a = ap.parse_args(argv)
    if not a.pdf.is_file():
        print(f"ERRO: arquivo não encontrado: {a.pdf}")
        return 2

    res = extrair(a.pdf)
    destino = a.saida or a.pdf.with_suffix(".xlsx")
    if not res.completo:
        destino = destino.with_name(destino.stem + "_INCOMPLETA.xlsx")
    gravar_xlsx(res, destino)
    for f in verificar_xlsx(res, destino):
        res.erros.append(f"Verificação do .xlsx falhou: {f}")

    soma = sum(p.quantidade or 0 for p in res.produtos)
    print(f"Arquivo: {destino}")
    print(f"Frete: {res.frete or 'não identificado'}")
    print(f"Produtos extraídos: {len(res.produtos)}" + (f" (PDF informa {res.total_produtos_doc})" if res.total_produtos_doc is not None else ""))
    print(f"Soma das unidades: {soma}" + (f" (PDF informa {res.total_unidades_doc})" if res.total_unidades_doc is not None else ""))
    for p in res.produtos:
        q = p.quantidade if p.quantidade is not None else "???"
        print(f"  {p.posicao:>2}. {p.codigo_ml or '???'} | {q} | {p.sku or '???'}")
    for av in res.avisos:
        print(f"AVISO: {av}")
    for er in res.erros:
        print(f"ERRO: {er}")
    if res.erros or not res.completo:
        print("STATUS: INCOMPLETO — NÃO usar como resultado validado.")
        return 1
    print("STATUS: OK — todos os campos confirmados e totais conferem." if res.total_unidades_doc is not None
          else "STATUS: OK (totais do PDF indisponíveis para conferência).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
