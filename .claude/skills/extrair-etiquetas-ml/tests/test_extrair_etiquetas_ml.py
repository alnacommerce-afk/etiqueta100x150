"""Testes da skill extrair-etiquetas-ml. Rodar: python -m unittest discover -s .claude/skills/extrair-etiquetas-ml/tests

O PDF de referência (Inbound-79427953) NÃO fica no Git: o teste reconstrói um PDF sintético com o mesmo
layout/coordenadas (inclusive as armadilhas: quantidades fora de ordem no texto bruto, números na descrição,
palavra "Unidades" minúscula, cabeçalho repetido na página 2). Se a variável ML_PDF_REFERENCIA apontar para o
PDF real, ele também é testado.
"""
import os
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
import pymupdf

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import extrair_etiquetas_ml as ex  # noqa: E402

ESPERADO = [
    ("EHAO37040", 25, "ROL_MAS_M"),
    ("ESPH23550", 20, "COLHER-MAD-26-5"),
    ("EZGH37086", 5, "TAB_COR_G"),
    ("GDHA38567", 10, "TAB_COR_G_2"),
    ("GXJW13905", 40, "COLHER-MAD-30-3"),
    ("JIBZ63695", 10, "MESA_DOB"),
    ("RNMD99644", 10, "TAB_COR_M"),
    ("XBUA71448", 5, "COLHER-MAD-30-5"),
    ("AKYR66331", 30, "COLHER-MAD-26-3"),
    ("IMIC32356", 30, "TAB_COR_G"),
    ("RABL13055", 40, "PILAO_G"),
    ("ZFWZ24473", 5, "COLHER-MAD-30-5"),
    ("ZLYK50731", 5, "COLHER-MAD-38-10"),
]
# descrições com números/palavras que imitam quantidade, como no PDF real
DESCR = {
    0: "Rolo De Massa Giratorio 42 Cm Para Pizza",
    3: "Kit 2 Tabuas De Corte Carne Madeira 2 Tabuas",
    4: "Kit 3 Colheres De Pau 30cm Concha 3 Colheres",
    12: "Kit Colher De Pau Atacado Revenda 5 A 15",
}


def _row(page, y, codigo, qtd, sku, descr, pular_qtd=False, total_extra=None):
    page.insert_text((38, y + 7), f"Código ML: {codigo} Código universal: N/A SKU:", fontsize=7.1)
    page.insert_text((38, y + 15), sku, fontsize=7.1)
    page.insert_text((38, y + 27), descr, fontsize=7.1)
    page.insert_text((38, y + 35), "Unidades extras 10", fontsize=7.1)
    page.insert_text((291, y + 7), "Etiquetagem", fontsize=7.1)
    page.insert_text((291, y + 15), "obrigatória", fontsize=7.1)
    if not pular_qtd:
        page.insert_text((237, y + 4), str(qtd), fontsize=7.1)


def montar_pdf(destino, linhas=ESPERADO, total_prod=13, total_un=235, pular_qtd=(), sku_vazio=()):
    doc = pymupdf.open()
    doc.new_page(width=595.28, height=841.89)
    doc.new_page(width=595.28, height=841.89)
    p1, p2 = doc[0], doc[1]
    p1.insert_text((40, 60), "Frete #79427953", fontsize=12)
    p1.insert_text((40, 72), f"Produtos do envio: {total_prod} | Total de unidades: {total_un}", fontsize=9)
    p1.insert_text((40, 130), "VERIFIQUE A EMBALAGEM E A IDENTIFICAÇÃO DOS SEUS PRODUTOS", fontsize=7)
    for x, t in ((38, "PRODUTO"), (237, "UNIDADES"), (291, "IDENTIFICAÇÃO"), (376, "INSTRUÇÕES"), (434, "PREPARAÇÃO")):
        p1.insert_text((x, 211), t, fontsize=7.1)
    for x, t in ((38, "PRODUTO"), (237, "UNIDADES"), (291, "IDENTIFICAÇÃO"), (376, "INSTRUÇÕES"), (434, "PREPARAÇÃO")):
        p2.insert_text((x, 60), t, fontsize=7.1)
    for i, (cod, q, sku) in enumerate(linhas):
        page, y = (p1, 233 + 61 * i) if i < 9 else (p2, 81 + 61 * (i - 9))
        _row(page, y - 7, cod, q, "" if i in sku_vazio else sku, DESCR.get(i, "Produto de teste"), pular_qtd=i in pular_qtd)
    p2.insert_text((428, 145), "plástico-bolha", fontsize=7.1)  # texto solto fora da coluna de unidades
    doc.save(destino)


class TestExtracao(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def _pdf(self, **kw):
        p = self.dir / "teste.pdf"
        montar_pdf(str(p), **kw)
        return p

    def test_referencia_completa(self):
        res = ex.extrair(self._pdf())
        self.assertEqual(res.erros, [])
        self.assertTrue(res.completo)
        self.assertEqual(len(res.produtos), 13)
        self.assertEqual(sum(p.quantidade for p in res.produtos), 235)
        self.assertEqual([(p.codigo_ml, p.quantidade, p.sku) for p in res.produtos], ESPERADO)
        self.assertEqual((res.total_produtos_doc, res.total_unidades_doc), (13, 235))

    def test_primeira_linha(self):
        res = ex.extrair(self._pdf())
        p = res.produtos[0]
        self.assertEqual((p.codigo_ml, p.quantidade, p.sku), ("EHAO37040", 25, "ROL_MAS_M"))

    def test_xlsx_gerado_e_verificado(self):
        pdf = self._pdf()
        out = self.dir / "saida.xlsx"
        rc = ex.main([str(pdf), "-o", str(out)])
        self.assertEqual(rc, 0)
        self.assertTrue(out.exists())
        self.assertEqual(ex.verificar_xlsx(ex.extrair(pdf), out), [])
        ws = openpyxl.load_workbook(out).worksheets[0]
        self.assertEqual([c.value for c in ws[1]], ["Código ML", "quantidade", "SKU"])
        self.assertEqual(ws.max_row, 14)
        self.assertEqual(ws.max_column, 3)
        self.assertEqual([c.value for c in ws[2]], ["EHAO37040", 25, "ROL_MAS_M"])
        self.assertIsInstance(ws["B2"].value, int)
        self.assertEqual(sum(ws.cell(r, 2).value for r in range(2, 15)), 235)

    def test_duplicados_sao_mantidos(self):  # TAB_COR_G e COLHER-MAD-30-5 aparecem 2x com códigos ML distintos
        res = ex.extrair(self._pdf())
        self.assertEqual([p.sku for p in res.produtos].count("TAB_COR_G"), 2)
        self.assertEqual([p.sku for p in res.produtos].count("COLHER-MAD-30-5"), 2)

    def test_quantidade_ausente_marca_incompleto(self):
        pdf = self._pdf(pular_qtd=(4,))
        res = ex.extrair(pdf)
        self.assertFalse(res.completo)
        self.assertIsNone(res.produtos[4].quantidade)
        self.assertTrue(any("quantidade ausente" in e and "GXJW13905" in e for e in res.erros))
        self.assertTrue(any("Divergência" in e for e in res.erros))  # soma != 235
        rc = ex.main([str(pdf), "-o", str(self.dir / "x.xlsx")])
        self.assertEqual(rc, 1)
        self.assertTrue((self.dir / "x_INCOMPLETA.xlsx").exists())
        self.assertFalse((self.dir / "x.xlsx").exists())

    def test_sku_ausente_marca_incompleto(self):
        res = ex.extrair(self._pdf(sku_vazio=(2,)))
        self.assertFalse(res.completo)
        self.assertTrue(any("SKU" in e for e in res.erros))

    def test_total_divergente_nao_e_ajustado(self):
        res = ex.extrair(self._pdf(total_un=240))
        self.assertFalse(res.completo)
        self.assertEqual(sum(p.quantidade for p in res.produtos), 235)  # soma real preservada
        self.assertTrue(any("240" in e for e in res.erros))

    def test_total_de_produtos_divergente(self):
        res = ex.extrair(self._pdf(total_prod=14))
        self.assertTrue(any("14 produtos" in e for e in res.erros))

    def test_pdf_inexistente(self):
        self.assertEqual(ex.main([str(self.dir / "nao_existe.pdf")]), 2)

    @unittest.skipUnless(os.environ.get("ML_PDF_REFERENCIA"), "defina ML_PDF_REFERENCIA para testar o PDF real")
    def test_pdf_real(self):
        res = ex.extrair(Path(os.environ["ML_PDF_REFERENCIA"]))
        self.assertEqual([(p.codigo_ml, p.quantidade, p.sku) for p in res.produtos], ESPERADO)
        self.assertTrue(res.completo)


if __name__ == "__main__":
    unittest.main()
